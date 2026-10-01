#!/usr/bin/env python3
from calendar import monthrange
from concurrent.futures import ThreadPoolExecutor
from datetime import datetime
from pathlib import Path

import cdsapi
import numpy as np
import pandas as pd
import xarray as xr
from openpyxl import load_workbook

INPUT = Path("SAMPLE_Fruit_morphology.xlsx")
DATA = Path("era5land_cds_2026")
OUTPUT = Path("SAMPLE_Fruit_morphology_ERA5Land_CDS_140d.xlsx")
VARIABLES = [
    "2m_temperature",
    "2m_dewpoint_temperature",
    "total_precipitation",
    "potential_evaporation",
    "surface_solar_radiation_downwards",
]
CORRECTIONS = {
    "26TAN-V3-C1": ("2026-07-12", 12.721722, 108.27825, "Date matched to V2; original coordinates"),
    "26TAN-V3-C2": ("2026-07-12", 12.721722, 108.27825, "Date matched to V2; coordinates from V3-C1"),
    "26TAN-V3-C3": ("2026-07-12", 12.721722, 108.27825, "Date matched to V2; coordinates from V3-C1"),
    "26TAN-V4-C1": ("2026-07-12", 12.725222, 108.276806, "Date matched to V2; original coordinates"),
    "26TAN-V4-C2": ("2026-07-12", 12.725222, 108.276806, "Date matched to V2; coordinates from V4-C1"),
    "26TAN-V4-C3": ("2026-07-12", 12.725222, 108.276806, "Date matched to V2; coordinates from V4-C1"),
    "26TAN-N1-V8-C2": ("2026-08-16", 12.633328, 108.151154, "Date matched to N1-V8-C1; coordinates midpoint C1-C3"),
}


def load_samples():
    ws = load_workbook(INPUT, read_only=True, data_only=True)["Samples_log"]
    header = [cell.value for cell in next(ws.iter_rows())]
    rows = []
    for values in ws.iter_rows(min_row=2, values_only=True):
        row = dict(zip(header, values))
        original_date = row["Ngày lấy"]
        correction = CORRECTIONS.get(row["Fruit-ID"])
        try:
            date = datetime.strptime(str(original_date), "%d %b %Y at %H:%M")
            note = "Original sample log"
        except (TypeError, ValueError):
            if not correction:
                continue
            date = datetime.fromisoformat(correction[0])
            row["Lat"], row["Long"], note = correction[1:]
        if row["Lat"] is not None and row["Long"] is not None:
            rows.append(
                {
                    **row,
                    "sampling_date": pd.Timestamp(date.date()),
                    "sampling_date_used": date.date(),
                    "correction_note": note,
                }
            )
    return pd.DataFrame(rows)


def download():
    DATA.mkdir(exist_ok=True)
    client = cdsapi.Client()
    for month in range(2, 5):
        target = DATA / f"era5land_2026_{month:02d}.nc"
        if target.exists() and target.stat().st_size > 10_000:
            continue
        first = 20 if month == 2 else 1
        last = 18 if month == 8 else monthrange(2026, month)[1]
        client.retrieve(
            "reanalysis-era5-land",
            {
                "variable": VARIABLES,
                "year": "2026",
                "month": f"{month:02d}",
                "day": [f"{day:02d}" for day in range(first, last + 1)],
                "time": [f"{hour:02d}:00" for hour in range(24)],
                "area": [14.1, 107.7, 12.0, 108.4],
                "data_format": "netcdf",
                "download_format": "unarchived",
            },
            str(target),
        )
    def fetch(month):
        target = DATA / f"era5land_2026_{month:02d}.nc"
        if target.exists() and target.stat().st_size > 10_000:
            return
        cdsapi.Client().retrieve(
                "reanalysis-era5-land",
                {
                    "variable": VARIABLES,
                    "year": "2026",
                    "month": f"{month:02d}",
                    "day": [f"{day:02d}" for day in range(1, 32)],
                    "time": [f"{hour:02d}:00" for hour in range(24)],
                    "area": [14.1, 107.7, 12.0, 108.4],
                    "data_format": "netcdf",
                    "download_format": "unarchived",
                },
                str(target),
            )
    with ThreadPoolExecutor(max_workers=4) as pool:
        list(pool.map(fetch, range(5, 9)))


def deaccumulate(series, clip=True):
    difference = series.diff("valid_time", label="upper").reindex(valid_time=series.valid_time)
    reset = series.valid_time.dt.hour == 1
    result = xr.where(difference.isnull() | reset, series, difference)
    return result.clip(min=0) if clip else result


def main():
    samples = load_samples()
    download()
    opened = [xr.open_dataset(path) for path in sorted(DATA.glob("*.nc"))]
    ds = xr.concat(opened, dim="valid_time").sortby("valid_time")
    ds = ds.sel(valid_time=~ds.get_index("valid_time").duplicated())
    times = pd.DatetimeIndex(ds.valid_time.values) + pd.Timedelta(hours=7)
    rain = deaccumulate(ds.tp) * 1000
    pev = deaccumulate(-ds.pev, clip=False) * 1000
    solar = deaccumulate(ds.ssrd) / 3600
    records = []

    for _, sample in samples.iterrows():
        point = {"latitude": sample.Lat, "longitude": sample.Long}
        grid = ds.sel(**point, method="nearest")
        hourly = pd.DataFrame(
            {
                "temperature_C": grid.t2m.values - 273.15,
                "dewpoint_C": grid.d2m.values - 273.15,
                "rain_mm": rain.sel(**point, method="nearest").values,
                "PEV_mm": pev.sel(**point, method="nearest").values,
                "solar_W_m2": solar.sel(**point, method="nearest").values,
            },
            index=times,
        )
        t, td = hourly.temperature_C, hourly.dewpoint_C
        hourly["RH_pct"] = np.clip(
            100 * np.exp(17.62 * td / (243.12 + td)) / np.exp(17.62 * t / (243.12 + t)), 0, 100
        )
        start = sample.sampling_date - pd.Timedelta(days=140)
        end = sample.sampling_date - pd.Timedelta(days=1)
        hourly = hourly.loc[str(start.date()) : str(end.date())]
        daily = hourly.resample("D").agg(
            temperature_mean_C=("temperature_C", "mean"),
            temperature_min_C=("temperature_C", "min"),
            temperature_max_C=("temperature_C", "max"),
            RH_mean_pct=("RH_pct", "mean"),
            rain_mm=("rain_mm", "sum"),
            PEV_mm=("PEV_mm", "sum"),
            sunshine_hours=("solar_W_m2", lambda values: (values >= 120).sum()),
        )
        for date, row in daily.iterrows():
            records.append({"Fruit-ID": sample["Fruit-ID"], "date": date, **row.to_dict()})

    daily = pd.DataFrame(records)
    summary = daily.groupby("Fruit-ID", as_index=False).agg(
        valid_days=("date", "count"),
        temperature_mean_C=("temperature_mean_C", "mean"),
        temperature_min_C=("temperature_min_C", "min"),
        temperature_max_C=("temperature_max_C", "max"),
        RH_mean_pct=("RH_mean_pct", "mean"),
        rain_total_mm=("rain_mm", "sum"),
        PEV_total_mm=("PEV_mm", "sum"),
        sunshine_total_hours=("sunshine_hours", "sum"),
    )
    meta = samples[
        ["Fruit-ID", "Ngày lấy", "sampling_date_used", "Lat", "Long", "correction_note"]
    ].copy()
    meta = meta.rename(
        columns={
            "Ngày lấy": "sampling_date_original",
            "sampling_date_used": "sampling_date_used",
            "correction_note": "data_note",
        }
    )
    meta["window_start"] = samples.sampling_date - pd.Timedelta(days=140)
    meta["window_end"] = samples.sampling_date - pd.Timedelta(days=1)
    summary = meta.merge(summary, on="Fruit-ID", how="left")
    summary[summary.select_dtypes("number").columns] = summary.select_dtypes("number").round(3)
    daily[daily.select_dtypes("number").columns] = daily.select_dtypes("number").round(3)

    with pd.ExcelWriter(OUTPUT, engine="openpyxl") as writer:
        summary.to_excel(writer, sheet_name="Summary_140d", index=False)
        daily.to_excel(writer, sheet_name="Daily_140d", index=False)
        pd.DataFrame(
            [
                ["Source", "Copernicus CDS reanalysis-era5-land, downloaded directly as NetCDF."],
                ["Window", "D-140 through D-1; sampling date excluded; daily grouping uses UTC+7."],
                ["Humidity", "Derived hourly from ERA5-Land T2m and D2m using the Magnus formula."],
                ["PEV", "Native ERA5-Land potential evaporation; sign reversed to positive mm."],
                ["Sunshine", "Estimated hours with ERA5-Land surface solar radiation >=120 W/m²."],
                ["Corrections", "7 missing dates supplied by user; 5 missing coordinates inferred from the same farm."],
            ],
            columns=["Item", "Definition"],
        ).to_excel(writer, sheet_name="README", index=False)

    assert len(summary) == 128 and summary.valid_days.eq(140).all()
    assert not summary.drop(columns="sampling_date_original").isna().any().any()
    print(f"Wrote {OUTPUT}: {len(summary)} samples, {len(daily)} daily rows")


if __name__ == "__main__":
    main()
