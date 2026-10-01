#!/usr/bin/env python3
from datetime import datetime
from pathlib import Path
import zipfile

import numpy as np
import pandas as pd
import xarray as xr
from openpyxl import load_workbook

INPUT = Path("DATA-2526CM-Pseudo-14Sep-La-Ce.xlsx")
RAW_DIR = Path("era5_monthly_presampling_10y")
OUTPUT = Path("DATA-2526CM-Pseudo-14Sep-La-Ce_ERA5_monthly_10y_presampling.xlsx")


def load_samples():
    ws = load_workbook(INPUT, read_only=True, data_only=True)["Sheet1"]
    header = [cell.value for cell in next(ws.iter_rows())]
    rows = []
    for values in ws.iter_rows(min_row=2, values_only=True):
        row = dict(zip(header, values))
        value = row["Ngày lấy"]
        date = value if isinstance(value, datetime) else datetime.strptime(str(value), "%d %b %Y at %H:%M")
        rows.append({**row, "sampling_date": date})
    return pd.DataFrame(rows)


def load_raw():
    datasets = []
    for archive in sorted(RAW_DIR.glob("*.nc")):
        if zipfile.is_zipfile(archive):
            parts = RAW_DIR / f"{archive.stem}_parts"
            parts.mkdir(exist_ok=True)
            with zipfile.ZipFile(archive) as zipped:
                zipped.extractall(parts)
            group = []
            for path in sorted(parts.glob("*.nc")):
                part = xr.open_dataset(path)
                month = pd.DatetimeIndex(part.valid_time.values).to_period("M").to_timestamp()
                group.append(part.assign_coords(valid_time=month))
            pev = xr.open_dataset(RAW_DIR / "pev_archives" / archive.name)
            pev_month = pd.DatetimeIndex(pev.valid_time.values).to_period("M").to_timestamp()
            group.append(pev.assign_coords(valid_time=pev_month))
            datasets.append(xr.merge(group, join="exact", compat="override"))
        else:
            datasets.append(xr.open_dataset(archive))
    return xr.concat(datasets, dim="valid_time").sortby("valid_time")


def main():
    samples = load_samples()
    if len(list(RAW_DIR.glob("*.nc"))) != 3:
        raise FileNotFoundError(f"Expected 3 CDS archives in {RAW_DIR}")
    ds = load_raw()
    times = pd.DatetimeIndex(ds.valid_time.values)
    monthly_rows = []

    for _, sample in samples.iterrows():
        point = ds.sel(latitude=sample.Lat, longitude=sample.Long, method="nearest")
        window_end = pd.Timestamp(sample.sampling_date).to_period("M") - 1
        window_start = window_end - 119
        keep = (times.to_period("M") >= window_start) & (times.to_period("M") <= window_end)
        point = point.isel(valid_time=keep)
        sample_times = times[keep]
        days = sample_times.days_in_month.to_numpy()
        temperature = point.t2m.values - 273.15
        dewpoint = point.d2m.values - 273.15
        humidity = np.clip(
            100 * np.exp(17.62 * dewpoint / (243.12 + dewpoint))
            / np.exp(17.62 * temperature / (243.12 + temperature)),
            0,
            100,
        )
        frame = pd.DataFrame(
            {
                "Sample-ID": sample["Sample-ID"],
                "Site": sample.Site,
                "sample_lat": sample.Lat,
                "sample_long": sample.Long,
                "grid_lat": float(point.latitude),
                "grid_long": float(point.longitude),
                "sampling_date": sample.sampling_date,
                "window_start": window_start.to_timestamp(),
                "window_end": window_end.to_timestamp(),
                "month": sample_times,
                "days": days,
                "temperature_C": temperature,
                "RH_proxy_pct": humidity,
                "wind_speed_10m_m_s": point.si10.values,
                "rainfall_mm_month": point.tp.values * 1000 * days,
                "PEV_mm_month": -point.pev.values * 1000 * days,
                "sunshine_equivalent_hours_month": point.ssrd.values * days / 3_600_000,
            }
        )
        monthly_rows.append(frame)

    monthly = pd.concat(monthly_rows, ignore_index=True)
    monthly[monthly.select_dtypes("number").columns] = monthly.select_dtypes("number").round(3)
    weighted = monthly.assign(
        temperature_x_days=monthly.temperature_C * monthly.days,
        humidity_x_days=monthly.RH_proxy_pct * monthly.days,
        wind_x_days=monthly.wind_speed_10m_m_s * monthly.days,
    ).groupby("Sample-ID", as_index=False).agg(
        window_start=("window_start", "first"),
        window_end=("window_end", "first"),
        valid_months=("month", "count"),
        days=("days", "sum"),
        temperature_x_days=("temperature_x_days", "sum"),
        humidity_x_days=("humidity_x_days", "sum"),
        wind_x_days=("wind_x_days", "sum"),
        rainfall_total_10y_mm=("rainfall_mm_month", "sum"),
        PEV_total_10y_mm=("PEV_mm_month", "sum"),
        sunshine_equivalent_total_10y_hours=("sunshine_equivalent_hours_month", "sum"),
    )
    weighted["temperature_mean_C"] = weighted.temperature_x_days / weighted.days
    weighted["RH_proxy_mean_pct"] = weighted.humidity_x_days / weighted.days
    weighted["wind_speed_10m_mean_m_s"] = weighted.wind_x_days / weighted.days
    weighted["rainfall_mean_annual_mm"] = weighted.rainfall_total_10y_mm / 10
    weighted["PEV_mean_annual_mm"] = weighted.PEV_total_10y_mm / 10
    weighted["sunshine_equivalent_mean_annual_hours"] = weighted.sunshine_equivalent_total_10y_hours / 10
    weighted = weighted.drop(columns=["temperature_x_days", "humidity_x_days", "wind_x_days"])
    meta = samples[["Sample-ID", "Tên mẫu", "Site", "Lat", "Long", "sampling_date"]]
    summary = meta.merge(weighted, on="Sample-ID", how="left")
    summary[summary.select_dtypes("number").columns] = summary.select_dtypes("number").round(3)

    readme = pd.DataFrame(
        [
            ["Source", "Copernicus CDS ERA5 monthly averaged reanalysis on single levels."],
            ["Period", "Per sample: 120 complete months immediately before the sampling month; sampling month excluded."],
            ["Temperature", "Monthly mean 2 m temperature, °C."],
            ["Humidity", "RH proxy (%) derived from monthly mean T2m and D2m; not mean observed RH."],
            ["Wind", "ERA5 monthly mean 10 m wind speed, m/s."],
            ["Rainfall", "Monthly total (mm) = ERA5 monthly mean daily accumulation × days in month."],
            ["PEV", "Native ERA5 potential evaporation, sign reversed to positive mm; total and annual mean reported."],
            ["Sunshine", "Equivalent full-sun hours = monthly SSRD energy / 1 kW/m²; not WMO sunshine duration."],
            ["Spatial", "Nearest ERA5 0.25° grid cell for each sample coordinate."],
        ],
        columns=["Item", "Definition"],
    )
    with pd.ExcelWriter(OUTPUT, engine="openpyxl") as writer:
        summary.to_excel(writer, sheet_name="Summary_10y", index=False)
        monthly.to_excel(writer, sheet_name="Monthly_10y", index=False)
        readme.to_excel(writer, sheet_name="README", index=False)

    assert len(summary) == 262 and len(monthly) == 262 * 120
    assert summary.valid_months.eq(120).all()
    assert not summary.isna().any().any()
    print(f"Wrote {OUTPUT}: {len(summary)} samples, {len(monthly)} sample-month rows")


if __name__ == "__main__":
    main()
